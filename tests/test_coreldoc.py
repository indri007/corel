import os
import io
import tempfile
import pytest
from PIL import Image, ImageDraw
import pytesseract
from docx import Document
import pypdf

from utils.word_to_pdf import word_to_pdf, check_libreoffice_installed
from utils.security_validator import sanitize_and_validate_file, MAX_FILE_SIZE_BYTES
from utils.cached_renderer import compute_sha256
from utils.pdf_manager import merge_pdfs
from utils.searchable_pdf import create_searchable_pdf

# =============================================================================
# TEST 1: word_to_pdf mengembalikan path PDF valid dari .docx
# =============================================================================
def test_word_to_pdf_conversion():
    if not check_libreoffice_installed():
        pytest.skip("LibreOffice ('soffice') tidak terpasang di sistem pengujian.")

    doc = Document()
    doc.add_heading("Judul Dokumen Uji Tesis", 0)
    doc.add_paragraph("Ini adalah isi pengujian konversi Word ke PDF dengan format asli.")

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        doc.save(tmp.name)
        docx_path = tmp.name

    try:
        pdf_path = word_to_pdf(docx_path)
        assert os.path.exists(pdf_path)
        assert pdf_path.endswith(".pdf")
        assert os.path.getsize(pdf_path) > 1000
    finally:
        if os.path.exists(docx_path):
            os.remove(docx_path)

# =============================================================================
# TEST 2: OCR dengan lang='ind' menghasilkan teks pada gambar contoh
# =============================================================================
def test_ocr_language_execution():
    img = Image.new("RGB", (450, 120), color="white")
    draw = ImageDraw.Draw(img)
    draw.text((25, 45), "DOKUMEN RESMI 2026", fill="black")

    installed_langs = pytesseract.get_languages()
    test_lang = "ind" if "ind" in installed_langs else "eng"

    extracted_text = pytesseract.image_to_string(img, lang=test_lang)
    assert len(extracted_text.strip()) > 0
    assert any(term in extracted_text.upper() for term in ["DOKUMEN", "RESMI", "2026"])

# =============================================================================
# TEST 3: Sanitasi nama berkas berbahaya (path traversal)
# =============================================================================
def test_file_sanitization():
    dangerous_filename = "../../../etc/passwd.pdf"
    dummy_pdf_bytes = b"%PDF-1.4\n%ValidPDFHeader\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF"

    is_valid, safe_name, msg = sanitize_and_validate_file(
        dummy_pdf_bytes,
        dangerous_filename,
        allowed_extensions=["pdf"]
    )
    assert is_valid is True
    assert ".." not in safe_name
    assert "/" not in safe_name
    assert safe_name.endswith(".pdf")

# =============================================================================
# TEST 4: Batas ukuran file menolak berkas > 50 MB
# =============================================================================
def test_file_size_limit_rejection():
    # 51 MB mock payload
    oversized_bytes = b"0" * (51 * 1024 * 1024)

    is_valid, safe_name, msg = sanitize_and_validate_file(
        oversized_bytes,
        "tesis_raksasa.pdf",
        allowed_extensions=["pdf"],
        max_size_bytes=MAX_FILE_SIZE_BYTES
    )
    assert is_valid is False
    assert "melebihi batas" in msg.lower()

# =============================================================================
# TEST 5: Caching / Determinisme fungsi hash
# =============================================================================
def test_deterministic_sha256_hash():
    sample_data = b"CorelDoc Studio Pro Verified Content"
    hash_1 = compute_sha256(sample_data)
    hash_2 = compute_sha256(sample_data)

    assert hash_1 == hash_2
    assert len(hash_1) == 64

# =============================================================================
# TEST 6 (P1): Merge 2 PDF menghasilkan jumlah halaman yang benar
# =============================================================================
def test_merge_two_pdfs():
    writer1 = pypdf.PdfWriter()
    writer1.add_blank_page(width=300, height=300)
    buf1 = io.BytesIO()
    writer1.write(buf1)
    pdf1_bytes = buf1.getvalue()

    writer2 = pypdf.PdfWriter()
    writer2.add_blank_page(width=300, height=300)
    writer2.add_blank_page(width=300, height=300)
    buf2 = io.BytesIO()
    writer2.write(buf2)
    pdf2_bytes = buf2.getvalue()

    merged_bytes = merge_pdfs([pdf1_bytes, pdf2_bytes])
    merged_reader = pypdf.PdfReader(io.BytesIO(merged_bytes))
    assert len(merged_reader.pages) == 3

# =============================================================================
# TEST 7 (P1): Searchable PDF mengandung teks yang bisa diekstrak oleh pypdf
# =============================================================================
def test_searchable_pdf_has_extractable_text():
    # Buat PDF 1-halaman berisi citra bitmap teks
    img = Image.new("RGB", (400, 100), color="white")
    draw = ImageDraw.Draw(img)
    draw.text((20, 35), "UNIVERSITAS INDONESIA", fill="black")

    img_buf = io.BytesIO()
    img.save(img_buf, format="PDF")
    input_pdf_bytes = img_buf.getvalue()

    # Jalankan pembuatan searchable PDF
    searchable_bytes = create_searchable_pdf(input_pdf_bytes, lang="eng", deskew=False)
    assert len(searchable_bytes) > 0

    reader = pypdf.PdfReader(io.BytesIO(searchable_bytes))
    extracted = reader.pages[0].extract_text()
    assert "UNIVERSITAS" in extracted.upper() or "INDONESIA" in extracted.upper()
