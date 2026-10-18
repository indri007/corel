import io
import os
import sys
import logging
import base64
import tempfile
import zipfile
from pathlib import Path
from PIL import Image
import pymupdf
import streamlit as st

# Setup Logging Terstruktur
from utils.logging_setup import setup_logging
setup_logging()
logger = logging.getLogger("coreldoc")

# Import Modul Internasionalisasi (i18n)
from utils.i18n import t

# Konversi DOCX → PDF
try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# Import utilitas inti
from utils.stamp_generator import generate_vector_stamp_svg, svg_to_png_bytes
from utils.image_enhancer import clean_scanned_document, enhance_photo_professional, trace_bitmap_to_svg
from utils.pdf_processor import (
    render_pdf_page_to_image,
    enhance_existing_pdf,
    create_professional_document_from_text
)
from utils.security_validator import sanitize_and_validate_file, MAX_FILE_SIZE_BYTES
from utils.cached_renderer import get_cached_pdf_page_image, get_cached_pdf_page_count, compute_sha256
from utils.searchable_pdf import make_searchable_pdf, ocr_langs_available
import utils.pdf_ops as pops
from utils.trace_advanced import batch_trace, trace_centerline_potrace, trace_color_vtracer, optimize_svg
import utils.redact as redact
import utils.sign as signlib
import utils.qr_verify as qrv
from utils.cdr_import import cdr_to_svg, cdr_info

# -----------------------------------------------------------------------------
# Konfigurasi Halaman Streamlit
# -----------------------------------------------------------------------------
st.set_page_config(
    page_title="CorelDoc Studio | Professional Document & Vector Suite",
    page_icon="🎨",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Inisialisasi State Bahasa
if "lang" not in st.session_state:
    st.session_state.lang = "id"

with st.sidebar:
    st.session_state.lang = st.selectbox(
        "🌐 " + t("language", st.session_state.lang),
        options=["id", "en"],
        format_func=lambda x: "Bahasa Indonesia" if x == "id" else "English (US)",
        index=["id", "en"].index(st.session_state.lang),
        key="lang_sel",
    )
    st.markdown("---")
    st.caption("CorelDoc Studio Pro v2.5")
    st.caption("Engine: PyMuPDF + LibreOffice + Tesseract + Poppler + Potrace")

# -----------------------------------------------------------------------------
# Custom CSS Styling (Material Design 3 — Light Theme)
# -----------------------------------------------------------------------------
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
    
    html, body, [class*="css"] {
        font-family: 'Inter', sans-serif;
    }

    .main-header {
        background: linear-gradient(135deg, #E0F2F1 0%, #E3F2FD 50%, #FFF3E0 100%);
        border: 1px solid #B2DFDB;
        border-radius: 20px;
        padding: 24px 28px;
        margin-bottom: 20px;
        box-shadow: 0 2px 12px rgba(0, 137, 123, 0.08);
    }

    .badge-pill {
        display: inline-block;
        background: #E0F2F1;
        color: #00695C;
        border: 1px solid #80CBC4;
        padding: 4px 14px;
        border-radius: 20px;
        font-size: 0.8rem;
        font-weight: 700;
        margin-right: 8px;
    }

    .badge-accent {
        display: inline-block;
        background: #E3F2FD;
        color: #1565C0;
        border: 1px solid #90CAF9;
        padding: 4px 14px;
        border-radius: 20px;
        font-size: 0.8rem;
        font-weight: 700;
    }

    .feature-card {
        background: #FFFFFF;
        border: 1px solid #E0E0E0;
        border-radius: 16px;
        padding: 20px;
        margin-bottom: 16px;
        box-shadow: 0 1px 6px rgba(0,0,0,0.06);
    }
</style>
""", unsafe_allow_html=True)

# -----------------------------------------------------------------------------
# Header Utama
# -----------------------------------------------------------------------------
_logo_path = os.path.join(os.path.dirname(__file__), "assets", "logo.jpg")
if os.path.exists(_logo_path):
    with open(_logo_path, "rb") as _f:
        _logo_b64 = base64.b64encode(_f.read()).decode()
    _logo_html = f'<img src="data:image/jpeg;base64,{_logo_b64}" style="width: 64px; height: 64px; border-radius: 14px; object-fit: cover; box-shadow: 0 2px 8px rgba(0,0,0,0.1); margin-right: 16px;" />'
else:
    _logo_html = ""

st.markdown(f"""
<div class="main-header">
    <div style="display: flex; align-items: center; justify-content: space-between; flex-wrap: wrap; gap: 12px;">
        <div style="display: flex; align-items: center;">
            {_logo_html}
            <div>
                <div style="margin-bottom: 6px;">
                    <span class="badge-pill">CORELDRAW-INSPIRED</span>
                    <span class="badge-accent">ENTERPRISE EDITION</span>
                </div>
                <h1 style="margin: 0; font-size: 2.0rem; font-weight: 800; color: #1C1B1F;">
                    CorelDoc Studio <span style="color: #00897B; font-size: 1.6rem;">Pro</span>
                </h1>
                <p style="margin: 4px 0 0 0; color: #49454F; font-size: 0.95rem;">
                    Platform Rekayasa Dokumen Korporat, Vektor Tracing, Redaksi Permanen, & Tanda Tangan Kriptografis PAdES.
                </p>
            </div>
        </div>
        <div>
            <a href="https://github.com/indri007/corel" target="_blank" style="text-decoration: none;">
                <button style="background: #00897B; color: #FFFFFF; border: none; padding: 10px 18px; border-radius: 10px; cursor: pointer; font-weight: 600; font-size: 0.85rem;">
                    ⭐ GitHub: indri007/corel
                </button>
            </a>
        </div>
    </div>
</div>
""", unsafe_allow_html=True)

# -----------------------------------------------------------------------------
# Navigasi Tab
# -----------------------------------------------------------------------------
tab_pdf, tab_img, tab_trace, tab_stamp, tab_catalog = st.tabs([
    "📄 Penyempurna PDF & Tulisan",
    "🖼️ Pembersih Scan & Gambar",
    "⚡ PowerTRACE (Bitmap → Vektor)",
    "🏵️ Pembuat Stempel Resmi",
    "📚 Katalog Ekosistem Corel"
])

# =============================================================================
# TAB 1: PENYEMPURNA PDF & TULISAN
# =============================================================================
with tab_pdf:
    st.subheader("Transformasi Dokumen & PDF ke Standar Korporat")
    st.caption("Menambahkan Kop Surat vektor, garis pemisah ganda, stempel resmi, watermark, nomor halaman, dan QR code validasi.")

    # Expander Fitur Lanjutan P1 & P2
    with st.expander("🔍 Buat Searchable PDF (OCR Layer Teks)", expanded=False):
        st.caption("Tambahkan layer teks ke PDF scan agar bisa dicari & di-copy.")
        pdf_in = st.file_uploader("Upload PDF untuk Searchable", type=["pdf"], key="spdf_in")
        lang_opts = ocr_langs_available() or ["ind+eng"]
        spdf_lang = st.selectbox(
            "Bahasa OCR",
            options=["ind+eng", "ind", "eng"] + [l for l in lang_opts if l not in {"ind+eng", "ind", "eng"}],
            index=0, key="spdf_lang",
        )
        spdf_deskew = st.checkbox("Auto-deskew halaman scan", value=True, key="spdf_deskew")
        spdf_force = st.checkbox("Paksa OCR ulang (abaikan layer teks lama)", value=False, key="spdf_force")

        if pdf_in and st.button("Buat Searchable PDF", key="spdf_btn"):
            with tempfile.TemporaryDirectory() as td:
                src = os.path.join(td, "in.pdf")
                out = os.path.join(td, "out.pdf")
                with open(src, "wb") as f:
                    f.write(pdf_in.read())
                with st.spinner("Memproses OCR... (bisa 1-3 menit per 10 halaman)"):
                    try:
                        make_searchable_pdf(src, out, lang=spdf_lang, deskew=spdf_deskew, force_ocr=spdf_force)
                        with open(out, "rb") as f:
                            st.download_button(
                                "⬇️ Download Searchable PDF",
                                f, file_name="searchable.pdf", mime="application/pdf",
                            )
                        st.success("Berhasil dibuat.")
                    except Exception as e:
                        st.error(f"Gagal: {e}")

    with st.expander("📚 Kelola PDF (Merge, Split, Rotate, Reorder, Delete, Compress)", expanded=False):
        mode = st.radio(
            "Operasi",
            ["Merge", "Split", "Extract", "Rotate", "Reorder", "Delete", "Compress"],
            horizontal=True, key="pdf_mode",
        )
        if mode == "Merge":
            files = st.file_uploader("Pilih 2+ PDF", type=["pdf"], accept_multiple_files=True, key="mg_in")
            if files and len(files) >= 2 and st.button("Merge", key="mg_btn"):
                with tempfile.TemporaryDirectory() as td:
                    paths = []
                    for i, f in enumerate(files):
                        p = os.path.join(td, f"in_{i}.pdf")
                        with open(p, "wb") as w: w.write(f.read())
                        paths.append(p)
                    out = os.path.join(td, "merged.pdf")
                    pops.merge_pdfs(paths, out)
                    with open(out, "rb") as r:
                        st.download_button("⬇️ Download", r, file_name="merged.pdf", mime="application/pdf")
        else:
            f_in = st.file_uploader("Pilih 1 PDF", type=["pdf"], key=f"in_{mode}")
            if f_in:
                with tempfile.TemporaryDirectory() as td:
                    src = os.path.join(td, "in.pdf")
                    with open(src, "wb") as w: w.write(f_in.read())
                    from pypdf import PdfReader
                    total = len(PdfReader(src).pages)
                    st.caption(f"Total halaman: {total}")
                    if mode == "Split":
                        chunk = st.number_input("Halaman per file", min_value=1, max_value=total, value=min(10, total))
                        if st.button("Split", key="sp_btn"):
                            outs = pops.split_pdf(src, td, int(chunk))
                            buf = io.BytesIO()
                            with zipfile.ZipFile(buf, "w") as z:
                                for p in outs: z.write(p, arcname=os.path.basename(p))
                            st.download_button("⬇️ Download ZIP", buf.getvalue(), file_name="split.zip", mime="application/zip")
                    elif mode == "Extract":
                        ranges = st.text_input("Rentang (contoh 1-3,5)", value=f"1-{min(3, total)}")
                        if st.button("Extract", key="ex_btn"):
                            out = os.path.join(td, "out.pdf")
                            pops.extract_pages(src, out, ranges)
                            with open(out, "rb") as r: st.download_button("⬇️ Download", r, file_name="extracted.pdf", mime="application/pdf")
                    elif mode == "Rotate":
                        deg = st.selectbox("Derajat", [90, 180, 270], index=0)
                        if st.button("Rotate", key="ro_btn"):
                            out = os.path.join(td, "out.pdf")
                            pops.rotate_pdf(src, out, deg)
                            with open(out, "rb") as r: st.download_button("⬇️ Download", r, file_name="rotated.pdf", mime="application/pdf")
                    elif mode == "Reorder":
                        order = st.text_input("Urutan baru (pisah koma)", value=",".join(str(i+1) for i in range(total)))
                        if st.button("Reorder", key="re_btn"):
                            out = os.path.join(td, "out.pdf")
                            pops.reorder_pdf(src, out, order)
                            with open(out, "rb") as r: st.download_button("⬇️ Download", r, file_name="reordered.pdf", mime="application/pdf")
                    elif mode == "Delete":
                        ranges = st.text_input("Halaman yang dihapus (contoh 2,5-7)", value="1")
                        if st.button("Delete", key="de_btn"):
                            out = os.path.join(td, "out.pdf")
                            pops.delete_pages(src, out, ranges)
                            with open(out, "rb") as r: st.download_button("⬇️ Download", r, file_name="deleted.pdf", mime="application/pdf")
                    elif mode == "Compress":
                        if st.button("Compress", key="co_btn"):
                            out = os.path.join(td, "out.pdf")
                            pops.compress_pdf(src, out)
                            st.caption(f"Sebelum: {os.path.getsize(src)//1024} KB → Sesudah: {os.path.getsize(out)//1024} KB")
                            with open(out, "rb") as r: st.download_button("⬇️ Download", r, file_name="compressed.pdf", mime="application/pdf")

    with st.expander("🔒 Redaksi PDF (Hapus Teks/Objek Permanen)", expanded=False):
        st.caption("Menghapus konten secara permanen — bukan sekadar kotak hitam di atasnya.")
        rf = st.file_uploader("Upload PDF untuk Redaksi", type=["pdf"], key="red_in")
        if rf:
            from pypdf import PdfReader
            with tempfile.TemporaryDirectory() as td:
                src = os.path.join(td, "in.pdf")
                with open(src, "wb") as w: w.write(rf.read())
                n = len(PdfReader(src).pages)
                pi = st.number_input("Halaman (0-based)", 0, max(0, n - 1), 0, key="red_pi")
                mode_red = st.radio("Mode Redaksi", ["Koordinat", "Cari Teks"], horizontal=True, key="red_mode")
                if mode_red == "Koordinat":
                    c1, c2, c3, c4 = st.columns(4)
                    x0 = c1.number_input("x0", 0.0, value=50.0)
                    y0 = c2.number_input("y0", 0.0, value=50.0)
                    x1 = c3.number_input("x1", 0.0, value=300.0)
                    y1 = c4.number_input("y1", 0.0, value=100.0)
                    if st.button("Redaksi Area", key="red_btn"):
                        out = os.path.join(td, "out.pdf")
                        redact.redact_rect(src, out, int(pi), (x0, y0, x1, y1))
                        with open(out, "rb") as r:
                            st.download_button("⬇️ Download Redacted PDF", r, file_name="redacted.pdf", mime="application/pdf")
                else:
                    q = st.text_input("Teks yang dicari untuk Diredaksi", key="red_q")
                    if q and st.button("Redaksi Teks", key="red_btn2"):
                        out = os.path.join(td, "out.pdf")
                        try:
                            redact.redact_text(src, out, q)
                            with open(out, "rb") as r:
                                st.download_button("⬇️ Download Redacted PDF", r, file_name="redacted.pdf", mime="application/pdf")
                        except ValueError as e:
                            st.warning(str(e))

    with st.expander("✍️ Tanda Tangan Digital Kriptografis (PAdES)", expanded=False):
        st.caption("Tanda tangan kriptografis berstandar PAdES — terotentikasi penuh, bukan sekadar gambar cap.")
        ts = st.tabs(["Tanda Tangani", "Verifikasi", "Buat Sertifikat Self-Signed"])
        with ts[0]:
            pdf_ts = st.file_uploader("PDF", type=["pdf"], key="ts_pdf")
            key_f = st.file_uploader("Private Key (.pem)", type=["pem"], key="ts_key")
            cert_f = st.file_uploader("Sertifikat (.pem)", type=["pem"], key="ts_cert")
            reason = st.text_input("Alasan", value="Disetujui secara digital", key="ts_reason")
            if pdf_ts and key_f and cert_f and st.button("Tanda Tangani", key="ts_btn"):
                with tempfile.TemporaryDirectory() as td:
                    p = os.path.join(td, "in.pdf"); open(p, "wb").write(pdf_ts.read())
                    k = os.path.join(td, "k.pem"); open(k, "wb").write(key_f.read())
                    c = os.path.join(td, "c.pem"); open(c, "wb").write(cert_f.read())
                    o = os.path.join(td, "out.pdf")
                    try:
                        signlib.sign_pdf(p, o, k, c, reason=reason)
                        with open(o, "rb") as r:
                            st.download_button("⬇️ Download Signed PDF", r, file_name="signed.pdf", mime="application/pdf")
                    except Exception as e:
                        st.error(f"Gagal: {e}")
        with ts[1]:
            vs = st.file_uploader("PDF bertanda tangan", type=["pdf"], key="vs_pdf")
            if vs and st.button("Verifikasi Tanda Tangan", key="vs_btn"):
                with tempfile.TemporaryDirectory() as td:
                    p = os.path.join(td, "in.pdf"); open(p, "wb").write(vs.read())
                    try:
                        res = signlib.verify_signature(p)
                        st.json(res)
                    except Exception as e:
                        st.error(f"Gagal: {e}")
        with ts[2]:
            cn = st.text_input("Common Name (CN)", value="CorelDoc Studio", key="cn")
            if st.button("Buat Sertifikat Self-Signed", key="gen_btn"):
                with tempfile.TemporaryDirectory() as td:
                    c = os.path.join(td, "cert.pem"); k = os.path.join(td, "key.pem")
                    signlib.generate_self_signed_cert(c, k, common_name=cn)
                    st.download_button("⬇️ Cert (.pem)", open(c, "rb").read(), file_name="cert.pem", mime="application/x-pem-file")
                    st.download_button("⬇️ Private Key (.pem)", open(k, "rb").read(), file_name="key.pem", mime="application/x-pem-file")

    with st.expander("📥 Impor CorelDRAW (.cdr)", expanded=False):
        st.caption("Konversi .cdr ke SVG/PNG/PDF via libcdr / TraceDraw / pylibcdr.")
        cf = st.file_uploader("Upload .cdr", type=["cdr"], key="cdr_in")
        if cf:
            with tempfile.TemporaryDirectory() as td:
                src = os.path.join(td, "in.cdr")
                with open(src, "wb") as w: w.write(cf.read())
                try:
                    import utils.cdr_import as cdr
                    info = cdr.cdr_info(src)
                    st.json(info)
                    fmt = st.radio("Format output", ["SVG", "PNG (300 DPI)", "PDF"], horizontal=True, key="cdr_fmt")
                    if st.button("Konversi .CDR", key="cdr_btn"):
                        with st.spinner("Mengonversi berkas CorelDRAW..."):
                            if fmt == "SVG":
                                out = os.path.join(td, "out.svg")
                                cdr.cdr_to_svg(src, out)
                                mime = "image/svg+xml"; name = "out.svg"
                            elif fmt == "PNG (300 DPI)":
                                out = os.path.join(td, "out.png")
                                cdr.cdr_to_png(src, out, dpi=300)
                                mime = "image/png"; name = "out.png"
                            else:
                                out = os.path.join(td, "out.pdf")
                                cdr.cdr_to_pdf(src, out)
                                mime = "application/pdf"; name = "out.pdf"
                            st.download_button(f"⬇️ Download {fmt}", open(out, "rb").read(), file_name=name, mime=mime)
                            if fmt == "SVG":
                                st.components.v1.html(open(out).read(), height=500, scrolling=True)
                except Exception as e:
                    st.error(f"Gagal: {e}")

    with st.expander("🎨 Editor Vektor (SVG)", expanded=False):
        st.caption("Muat SVG, edit path, tambah shape, ekspor ulang.")
        ef = st.file_uploader("Upload SVG", type=["svg"], key="svg_ed_in")
        if ef:
            import utils.svg_editor as sve
            with tempfile.TemporaryDirectory() as td:
                src = os.path.join(td, "in.svg")
                with open(src, "wb") as w: w.write(ef.read())
                root = sve.load_svg(src)
                paths = sve.list_paths(root)
                st.caption(f"{len(paths)} path terdeteksi")

                if paths:
                    idx = st.selectbox("Pilih path", range(len(paths)),
                                       format_func=lambda i: f"#{i} id={paths[i]['id']}",
                                       key="svg_ed_idx")
                    c1, c2, c3 = st.columns(3)
                    new_fill = c1.color_picker("Fill", paths[idx]["fill"] if paths[idx]["fill"].startswith("#") else "#C8102E")
                    new_stroke = c2.color_picker("Stroke", "#000000")
                    new_sw = c3.slider("Stroke width", 0.0, 10.0, 1.0, 0.5)
                    dx = st.slider("Geser X", -100.0, 100.0, 0.0, 1.0)
                    dy = st.slider("Geser Y", -100.0, 100.0, 0.0, 1.0)
                    if st.button("Terapkan Perubahan Path", key="svg_ed_apply"):
                        sve.set_path_attr(root, idx, "fill", new_fill)
                        sve.set_path_attr(root, idx, "stroke", new_stroke)
                        sve.set_path_attr(root, idx, "stroke-width", str(new_sw))
                        if dx or dy:
                            sve.translate_path(root, idx, dx, dy)
                        st.session_state["edited_svg"] = root

                with st.form("add_shape", clear_on_submit=False):
                    st.write("**Tambah Shape Vektor Baru**")
                    sc1, sc2 = st.columns(2)
                    shape = sc1.selectbox("Jenis", ["Rectangle", "Circle", "Text"])
                    if shape == "Text":
                        txt = sc2.text_input("Teks", "Halo")
                        if st.form_submit_button("Tambah"):
                            sve.add_text(root, 50, 100, txt); st.session_state["edited_svg"] = root
                    elif shape == "Rectangle":
                        if st.form_submit_button("Tambah"):
                            sve.add_rect(root, 50, 50, 200, 100); st.session_state["edited_svg"] = root
                    else:
                        if st.form_submit_button("Tambah"):
                            sve.add_circle(root, 150, 150, 80); st.session_state["edited_svg"] = root

                if "edited_svg" in st.session_state:
                    out = os.path.join(td, "out.svg")
                    sve.save_svg(st.session_state["edited_svg"], out)
                    svg_text = open(out).read()
                    st.components.v1.html(svg_text, height=400, scrolling=True)
                    st.download_button("⬇️ Download SVG", svg_text, file_name="edited.svg", mime="image/svg+xml")

    with st.expander("🤖 AI: Text → SVG", expanded=False):
        st.caption("Generasi SVG dari deskripsi teks via Hugging Face Inference API.")
        import utils.ai_svg as aisvg
        if not os.environ.get("HF_TOKEN"):
            st.warning("Set `HF_TOKEN` di environment: `export HF_TOKEN=hf_xxx`")
        prompt = st.text_area("Deskripsi vektor",
                              "A minimalist red circular logo with white letter C in the center",
                              key="ai_prompt")
        model = st.selectbox("Model", [
            "Caffin/SVGThinker-7B",
            "starvector/starvector-8b-im2svg",
            "xingxm/HiVG-3B-Base",
        ], key="ai_model")
        if st.button("Generate SVG", key="ai_btn"):
            with st.spinner("Memanggil HF Inference API..."):
                try:
                    raw = aisvg.text_to_svg_hf(prompt, model=model)
                    svg_text = aisvg.extract_svg(raw)
                    with tempfile.TemporaryDirectory() as td:
                        p = aisvg.save_svg_text(svg_text, os.path.join(td, "out.svg"))
                        st.components.v1.html(svg_text, height=400, scrolling=True)
                        st.download_button("⬇️ Download SVG", svg_text, file_name="ai_generated.svg", mime="image/svg+xml")
                except Exception as e:
                    st.error(f"Gagal: {e}")

    with st.expander("📚 Batch Tesis/Buku (175 halaman)", expanded=False):
        st.caption("PDF tebal → split per bab + OCR searchable + trace halaman kunci → ZIP.")
        import utils.batch_pipeline as bp
        bf = st.file_uploader("Upload PDF tebal (Tesis / Buku 175+ halaman)", type=["pdf"], key="bp_in")
        if bf:
            from pypdf import PdfReader
            with tempfile.TemporaryDirectory() as td:
                src = os.path.join(td, "in.pdf")
                with open(src, "wb") as w: w.write(bf.read())
                total = len(PdfReader(src).pages)
                st.caption(f"Total halaman: **{total}**")

                c1, c2 = st.columns(2)
                chunk = c1.number_input("Halaman per bab", 5, 100, 20, key="bp_chunk")
                do_ocr = c2.checkbox("Searchable PDF per chunk", value=True, key="bp_ocr")
                trace_str = st.text_input("Halaman untuk di-trace (0-based, koma)", value="0",
                                          key="bp_trace",
                                          help="Contoh: 0,5,10 untuk cover + 2 bagan")
                cl = st.checkbox("Centerline tracing (line art)", value=False, key="bp_cl")

                if st.button("Jalankan Pipeline Tesis", key="bp_btn"):
                    trace_pages = [int(x) for x in trace_str.split(",") if x.strip().isdigit()]
                    bar = st.progress(0.0, text="Memulai...")
                    def _p(stage: str, cur: int, tot: int):
                        bar.progress(min(cur / max(tot, 1), 1.0), text=f"{stage} ({cur}/{tot})")
                    out_zip = os.path.join(td, "result.zip")
                    try:
                        bp.process_thesis(src, out_zip, chunk=int(chunk),
                                          do_searchable=do_ocr,
                                          trace_pages=trace_pages,
                                          trace_centerline=cl,
                                          progress=_p)
                        bar.progress(1.0, text="Selesai")
                        st.download_button("⬇️ Download Paket Tesis ZIP", open(out_zip, "rb").read(),
                                           file_name="thesis_batch.zip", mime="application/zip")
                    except Exception as e:
                        st.error(f"Gagal: {e}")

    st.markdown("---")

    col_input, col_preview = st.columns([1, 1.2])

    with col_input:
        input_type = st.radio(
            "Pilih Sumber Dokumen:",
            ["Unggah Berkas PDF / Word / CDR", "Tulis / Tempel Teks Baru"],
            horizontal=True
        )

        uploaded_pdf = None
        doc_title = ""
        doc_content = ""

        if input_type.startswith("Unggah"):
            uploaded_pdf = st.file_uploader(
                "Pilih Berkas PDF, Word (.docx/.doc), atau Corel (.cdr):",
                type=["pdf", "docx", "doc", "odt", "rtf", "cdr"],
                key="pdf_uploader",
                help="Maksimal 50 MB. Mendukung konversi otomatis Word (LibreOffice Engine) dan Corel (.cdr)."
            )
        else:
            doc_title = st.text_input("Judul Dokumen:", value="SURAT KEPUTUSAN DIREKSI TENTANG STANDARISASI DOKUMEN")
            doc_number = st.text_input("Nomor Registrasi / Surat:", value="SK/DIR-GTI/X/2026-088")
            doc_content = st.text_area(
                "Isi Naskah / Ketetapan Dokumen:",
                value="""Menimbang bahwa untuk mewujudkan tata kelola persuratan yang terintegrasi, profesional, dan akuntabel di lingkungan korporasi modern...

MEMUTUSKAN:
1. Menetapkan format dokumen korporat standar dengan identitas vektor resmi.
2. Mengesahkan penggunaan stempel digital dan verifikasi hash sha-256.
3. Seluruh naskah wajib memenuhi kaidah keterbacaan tinggi dan tata letak presisi.""",
                height=180
            )

        st.markdown("---")
        st.markdown("#### ⚙️ Pengaturan Desain & Ornamen")

        c1, c2 = st.columns(2)
        with c1:
            header_name = st.text_input("Nama Instansi / Perusahaan:", value="PT GLOBAL TEKNOLOGI INOVASI")
            accent_col = st.color_picker("Warna Aksen Korporat:", value="#00897B")
            watermark_val = st.text_input("Teks Watermark:", value="CONFIDENTIAL")
        with c2:
            header_sub = st.text_input("Sub-Judul / Divisi:", value="Directorate of Technology & Engineering Systems")
            footer_val = st.text_input("Teks Footer / Legal Disclaimer:", value="Dokumen Resmi Terverifikasi - Rahasia & Hak Cipta Terlindungi")

        st.markdown("##### Fitur Ornamen Vektor:")
        f1, f2, f3, f4, f5 = st.columns(5)
        with f1: en_header = st.checkbox("Kop Surat", value=True)
        with f2: en_footer = st.checkbox("Footer & Hal", value=True)
        with f3: en_border = st.checkbox("Border Ganda", value=True)
        with f4: en_watermark = st.checkbox("Watermark", value=True)
        with f5: en_stamp = st.checkbox("Stempel Resmi", value=True)

    with col_preview:
        st.markdown("#### 👁️ Pratinjau Dokumen Profesional")

        result_pdf_bytes = None

        if input_type.startswith("Unggah") and uploaded_pdf is not None:
            raw_bytes = uploaded_pdf.read()
            file_name_lower = uploaded_pdf.name.lower()

            # Sanitasi & Validasi Ukuran (Maks 50MB)
            valid, safe_fname, v_msg = sanitize_and_validate_file(
                raw_bytes, uploaded_pdf.name, allowed_extensions=["pdf", "docx", "doc", "odt", "rtf", "cdr"]
            )
            if not valid:
                st.error(f"❌ {v_msg}")
                st.stop()

            # Konversi .cdr (CorelDRAW) jika diunggah
            if file_name_lower.endswith(".cdr"):
                with st.spinner("Mengonversi berkas CorelDRAW (.cdr) ke PDF/SVG..."):
                    with tempfile.TemporaryDirectory() as td:
                        cdr_in = os.path.join(td, "in.cdr")
                        pdf_out = os.path.join(td, "out.pdf")
                        with open(cdr_in, "wb") as f_c: f_c.write(raw_bytes)
                        try:
                            from utils.cdr_import import cdr_to_pdf
                            cdr_to_pdf(cdr_in, pdf_out)
                            with open(pdf_out, "rb") as f_p: raw_bytes = f_p.read()
                            st.success("✅ File CorelDRAW (.cdr) berhasil dikonversi ke PDF!")
                        except Exception as e_cdr:
                            st.error(f"Gagal mengimpor file CDR: {e_cdr}")
                            st.stop()

            # Konversi DOCX / Word → PDF dengan Full Fidelity via LibreOffice Engine
            elif file_name_lower.endswith((".docx", ".doc", ".odt", ".rtf")):
                with st.spinner("Mengonversi dokumen Word ke PDF dengan presisi tata letak penuh (LibreOffice Engine)..."):
                    try:
                        from utils.word_to_pdf import word_to_pdf
                        suffix = os.path.splitext(uploaded_pdf.name)[1]
                        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp_word:
                            tmp_word.write(raw_bytes)
                            tmp_word_path = tmp_word.name

                        converted_pdf_path = word_to_pdf(tmp_word_path)
                        with open(converted_pdf_path, "rb") as f_pdf:
                            raw_bytes = f_pdf.read()
                        st.success(f"✅ Dokumen Word berhasil dikonversi dengan tata letak asli utuh! ({len(raw_bytes):,} bytes)")
                    except Exception as err_lo:
                        st.warning(f"Engine LibreOffice headless tidak aktif ({err_lo}). Beralih ke ekstraksi teks ReportLab...")
                        if not HAS_DOCX:
                            st.error("Library `python-docx` tidak tersedia.")
                            st.stop()
                        try:
                            docx_doc = DocxDocument(io.BytesIO(raw_bytes))
                            full_text = "\n\n".join(
                                para.text for para in docx_doc.paragraphs if para.text.strip()
                            )
                            raw_bytes = create_professional_document_from_text(
                                title=doc_title if doc_title else uploaded_pdf.name.rsplit(".", 1)[0],
                                content=full_text,
                                company_name=header_name,
                                doc_number="DOCX-CONVERT",
                                accent_color_hex=accent_col
                            )
                        except Exception as e_docx:
                            st.error(f"Gagal memproses dokumen Word: {e_docx}")
                            st.stop()

            with st.spinner("Memproses transformasi dokumen..."):
                result_pdf_bytes = enhance_existing_pdf(
                    raw_bytes,
                    header_title=header_name,
                    header_subtitle=header_sub,
                    watermark_text=watermark_val,
                    footer_text=footer_val,
                    accent_color=accent_col,
                    apply_watermark=en_watermark,
                    apply_header=en_header,
                    apply_footer=en_footer,
                    apply_border=en_border,
                    apply_stamp=en_stamp
                )
        elif not input_type.startswith("Unggah") and doc_content.strip():
            with st.spinner("Menyusun dokumen profesional dari teks..."):
                result_pdf_bytes = create_professional_document_from_text(
                    title=doc_title,
                    content=doc_content,
                    company_name=header_name,
                    doc_number=doc_number,
                    accent_color_hex=accent_col
                )

        if result_pdf_bytes is not None:
            total_pages = get_cached_pdf_page_count(result_pdf_bytes)
            sha_hash = compute_sha256(result_pdf_bytes)

            st.success(f"✅ Dokumen berhasil diproses ({total_pages} Halaman)!")
            st.caption(f"🔒 SHA-256 Audit Hash: `{sha_hash}`")

            page_to_show = 1
            if total_pages > 1:
                page_to_show = st.slider("Pratinjau Halaman:", min_value=1, max_value=total_pages, value=1)

            preview_img = get_cached_pdf_page_image(result_pdf_bytes, page_number=page_to_show, dpi=140)
            st.image(preview_img, caption=f"Pratinjau Halaman {page_to_show} dari {total_pages}", use_container_width=True)

            st.download_button(
                label="📥 Unduh Dokumen PDF Resmi",
                data=result_pdf_bytes,
                file_name="dokumen_korporat_resmi.pdf",
                mime="application/pdf",
                use_container_width=True
            )
        else:
            st.info("💡 Unggah berkas dokumen atau ketik naskah di sebelah kiri untuk melihat hasil pratinjau.")

# =============================================================================
# TAB 2: PEMBERSIH SCAN & GAMBAR
# =============================================================================
with tab_img:
    st.subheader("Penyempurnaan Gambar & Pembersihan Scan Dokumen")
    st.caption("Memutihkan latar belakang kertas yang buram/kotor, menajamkan teks tulisan, serta mengoptimalkan kontras & vibrance gambar.")

    col_img_in, col_img_out = st.columns([1, 1.2])

    with col_img_in:
        uploaded_img = st.file_uploader(
            "Unggah Gambar, Scan Dokumen, atau Berkas Word/PDF:",
            type=["png", "jpg", "jpeg", "webp", "docx", "doc", "pdf"],
            key="img_enhancer_uploader",
            help="Mendukung format gambar (PNG/JPG), dokumen Word (.docx), dan PDF. Maksimal 50 MB."
        )

        mode_process = st.selectbox(
            "Pilih Mode Pemrosesan:",
            ["Pembersih Scan Dokumen (White Background & Sharp Text)", "Penyempurna Foto & Ilustrasi (Tone, Vibrance, Sharpness)"]
        )

        if mode_process.startswith("Pembersih"):
            aggressiveness = st.slider("Tingkat Agresivitas Pemutihan Kertas:", 0.5, 2.0, 1.2, 0.1)
        else:
            b_val = st.slider("Brightness:", 0.5, 2.0, 1.05, 0.05)
            c_val = st.slider("Contrast:", 0.5, 2.0, 1.2, 0.05)
            s_val = st.slider("Sharpness:", 0.5, 3.0, 1.5, 0.1)
            col_val = st.slider("Color Vibrance:", 0.0, 2.0, 1.15, 0.05)

        # Opsi Bahasa OCR
        installed_ocr_langs = ocr_langs_available() or ["ind", "eng"]
        ocr_lang_choice = st.selectbox(
            "Bahasa Mesin OCR (Tesseract):",
            options=["ind+eng", "ind", "eng"] + [l for l in installed_ocr_langs if l not in {"ind+eng", "ind", "eng"}],
            index=0, key="ocr_lang_choice"
        )

    with col_img_out:
        if uploaded_img is not None:
            raw_bytes = uploaded_img.read()
            uploaded_name_lower = uploaded_img.name.lower()

            # Sanitasi & Validasi
            valid, safe_fname, v_msg = sanitize_and_validate_file(
                raw_bytes, uploaded_img.name, allowed_extensions=["png", "jpg", "jpeg", "webp", "docx", "doc", "pdf"]
            )
            if not valid:
                st.error(f"❌ {v_msg}")
                st.stop()

            # Deteksi jika yang diunggah adalah Dokumen Word (.docx)
            if uploaded_name_lower.endswith((".docx", ".doc")):
                st.info(f"📄 Berkas Word terdeteksi: **{uploaded_img.name}**")
                if not HAS_DOCX:
                    st.error("Library `python-docx` belum terpasang.")
                    st.stop()
                
                try:
                    from docx import Document as DocxDoc
                    doc_obj = DocxDoc(io.BytesIO(raw_bytes))
                    full_text = "\n\n".join([p.text for p in doc_obj.paragraphs if p.text.strip()])
                    st.success(f"✅ Dokumen Word berhasil dibaca ({len(full_text):,} karakter teks ditemukan)!")
                    st.text_area("Isi Teks dari Berkas Word:", value=full_text, height=220)
                    st.stop()
                except Exception as e:
                    st.error(f"Gagal membaca berkas Word: {e}")
                    st.stop()

            # Deteksi jika yang diunggah adalah PDF
            elif uploaded_name_lower.endswith(".pdf"):
                try:
                    pdf_doc = pymupdf.open(stream=raw_bytes, filetype="pdf")
                    page_num = st.number_input("Pilih Halaman PDF untuk Diproses:", min_value=1, max_value=len(pdf_doc), value=1)
                    page = pdf_doc[page_num - 1]
                    pix = page.get_pixmap(dpi=150)
                    raw_img = Image.open(io.BytesIO(pix.tobytes("png")))
                except Exception as e:
                    st.error(f"Gagal membaca berkas PDF: {e}")
                    st.stop()
            else:
                raw_img = Image.open(io.BytesIO(raw_bytes))

            if mode_process.startswith("Pembersih"):
                processed_img = clean_scanned_document(raw_img, aggressiveness=aggressiveness)
            else:
                processed_img = enhance_photo_professional(raw_img, brightness=b_val, contrast=c_val, sharpness=s_val, color=col_val)

            c_orig, c_proc = st.columns(2)
            with c_orig:
                st.markdown("**Sebelum (Original):**")
                st.image(raw_img, use_container_width=True)
            with c_proc:
                st.markdown("**Sesudah (Profesional):**")
                st.image(processed_img, use_container_width=True)

            out_buf = io.BytesIO()
            processed_img.save(out_buf, format="PNG")
            st.download_button(
                label="📥 Unduh Gambar Berkualitas Tinggi (PNG)",
                data=out_buf.getvalue(),
                file_name="gambar_profesional.png",
                mime="image/png",
                use_container_width=True
            )

            # Fitur Ekstraksi Tulisan OCR Tesseract dengan Parameter Bahasa
            st.markdown("---")
            st.markdown("##### 🔍 Ekstraksi Tulisan (OCR Tesseract)")
            if st.button("Jalankan Pembacaan Teks (OCR)", key="btn_ocr"):
                with st.spinner(f"Membaca teks menggunakan Tesseract (Bahasa: {ocr_lang_choice})..."):
                    try:
                        import pytesseract
                        extracted_text = pytesseract.image_to_string(processed_img, lang=ocr_lang_choice)
                        if extracted_text.strip():
                            st.text_area("Hasil Ekstraksi Teks (Dapat Diedit/Disalin):", value=extracted_text, height=140)
                            st.success("Teks berhasil diekstrak! Anda dapat menyalin teks ini ke Tab 1 (Penyempurna PDF) untuk diformat ulang.")
                        else:
                            st.warning("Tidak ditemukan teks yang dapat dibaca secara jelas pada gambar ini.")
                    except Exception as e:
                        logger.exception("Gagal memproses OCR")
                        st.error(f"Gagal memproses OCR: {str(e)}")
        else:
            st.info("💡 Unggah gambar tanda tangan, diagram, scan kertas, atau foto untuk melakukan pembersihan & penajaman otomatis.")

# =============================================================================
# TAB 3: POWERTRACE (BITMAP TO VECTOR SVG)
# =============================================================================
with tab_trace:
    st.subheader("⚡ PowerTRACE: Vektorisasi Bitmap ke SVG Murni (Tingkat Produksi)")
    st.caption("Mengonversi logo, gambar, tanda tangan, atau sketsa bitmap menjadi garis vektor SVG berpresisi tinggi.")

    trace_mode = st.radio("Mode Tracing", ["Warna (VTracer)", "Centerline (Potrace)"], horizontal=True, key="tr_mode")
    centerline = trace_mode.startswith("Centerline")
    do_bg_remove = st.checkbox("Hapus background putih (Transparan)", value=True, key="tr_bg")
    do_svgo = st.checkbox("Optimasi SVG (SVGO jika terpasang)", value=True, key="tr_svgo")

    multi = st.file_uploader(
        "Upload 1 atau Lebih Gambar untuk Batch Tracing:",
        type=["png", "jpg", "jpeg", "webp"],
        accept_multiple_files=True,
        key="tr_files"
    )

    if multi and st.button("Jalankan PowerTRACE", key="tr_btn"):
        with tempfile.TemporaryDirectory() as td:
            paths = []
            for i, f in enumerate(multi):
                p = os.path.join(td, f"in_{i}_{f.name}")
                with open(p, "wb") as w: w.write(f.read())
                paths.append(p)
            bar = st.progress(0.0, text="Memproses vektorisasi...")
            def _p(i, n): bar.progress(i / n, text=f"Memproses gambar {i} dari {n}")
            outs = batch_trace(paths, td, centerline=centerline, remove_bg=do_bg_remove, do_svgo=do_svgo, progress=_p)
            st.success(f"✅ Berhasil memvektorisasi {len(outs)} gambar!")
            if len(outs) == 1:
                with open(outs[0], "rb") as r:
                    st.download_button("⬇️ Download SVG Vektor", r, file_name=os.path.basename(outs[0]), mime="image/svg+xml")
            else:
                buf = io.BytesIO()
                with zipfile.ZipFile(buf, "w") as z:
                    for p in outs: z.write(p, arcname=os.path.basename(p))
                st.download_button("⬇️ Download Batch SVG (ZIP)", buf.getvalue(), file_name="trace_batch.zip", mime="application/zip")

# =============================================================================
# TAB 4: PEMBUAT STEMPEL RESMI & VERIFIKASI QR
# =============================================================================
with tab_stamp:
    st.subheader("🏵️ Pembuat Stempel Resmi Vektor & Audit QR Verifikasi")
    st.caption("Menghasilkan stempel legal melingkar resolusi tinggi dan manifest verifikasi hash SHA-256 untuk otentikasi dokumen.")

    col_st_cfg, col_st_view = st.columns([1, 1.2])

    with col_st_cfg:
        s_org = st.text_input("Teks Melingkar Atas (Nama Organisasi):", value="KEMENTERIAN RISET & TEKNOLOGI")
        s_dept = st.text_input("Teks Melingkar Bawah (Unit Kerja):", value="DIREKTORAT JENDERAL INOVASI")
        s_status = st.text_input("Teks Kotak Tengah (Status Dokumen):", value="DISETUJUI & SAH")
        s_date = st.text_input("Tanggal Verifikasi:", value="18 OKT 2026")
        s_color = st.selectbox(
            "Warna Tinta Stempel:",
            ["#00897B (Teal Korporat)", "#0d47a1 (Biru Dinas)", "#b71c1c (Merah Resmi)", "#4a148c (Ungu Royal)"]
        )
        chosen_hex = s_color.split(" ")[0]

        st.markdown("---")
        st.markdown("##### 🏷️ Pengaturan URL Audit QR")
        base_url = st.text_input("Base URL Verifikasi:", value="https://coreldoc.example.com", key="qr_base")
        doc_id = st.text_input("ID Dokumen (UUID / Nomor Surat):", value="DOC-2026-0001", key="qr_id")
        pdf_for_qr = st.file_uploader("Upload PDF untuk Dihasilkan Manifest & QR:", type=["pdf"], key="qr_pdf")

    with col_st_view:
        stamp_svg = generate_vector_stamp_svg(
            org_name=s_org,
            dept_name=s_dept,
            status_text=s_status,
            date_str=s_date,
            color=chosen_hex,
            diameter=320
        )

        st.html(
            f'<div style="display: flex; justify-content: center; align-items: center; padding: 20px; background: #ffffff; border-radius: 12px; border: 1px solid #e0e0e0;">{stamp_svg}</div>'
        )

        c_d1, c_d2 = st.columns(2)
        with c_d1:
            st.download_button(
                label="📥 Unduh Stempel SVG",
                data=stamp_svg.encode("utf-8"),
                file_name="stempel_resmi.svg",
                mime="image/svg+xml",
                use_container_width=True
            )
        with c_d2:
            stamp_png_bytes = svg_to_png_bytes(stamp_svg, zoom=2)
            st.download_button(
                label="📥 Unduh Stempel PNG Transparan",
                data=stamp_png_bytes,
                file_name="stempel_resmi.png",
                mime="image/png",
                use_container_width=True
            )

        if pdf_for_qr:
            with tempfile.TemporaryDirectory() as td:
                p = os.path.join(td, "d.pdf")
                with open(p, "wb") as w: w.write(pdf_for_qr.read())
                man = qrv.make_manifest(p, doc_id, base_url)
                st.markdown("---")
                st.markdown("##### 📌 QR Verifikasi URL & SHA-256")
                st.code(man["url"])
                st.caption(f"SHA-256 Audit: `{man['sha256']}`")
                qr_file = qrv.make_qr(man["url"], os.path.join(td, "qr.png"))
                st.image(str(qr_file), width=180)
                
                cq1, cq2 = st.columns(2)
                with cq1:
                    st.download_button("⬇️ Download QR PNG", open(qr_file, "rb").read(), file_name="qr.png", mime="image/png")
                with cq2:
                    mj = qrv.write_manifest(man, os.path.join(td, "manifest.json"))
                    st.download_button("⬇️ Download Manifest JSON", open(mj, "rb").read(), file_name="manifest.json", mime="application/json")

# =============================================================================
# TAB 5: KATALOG EKOSISTEM COREL
# =============================================================================
with tab_catalog:
    st.subheader("📚 Katalog Ekosistem Corel & Alternatif Grafis Terbuka")
    st.caption("Referensi arsitektur format berkas CorelDRAW (.CDR) dan teknologi open-source pendukung.")

    st.markdown("""
| Tipe | Nama Perangkat Lunak | Bahasa Utama | Deskripsi & Kompatibilitas |
| :--- | :--- | :--- | :--- |
| **Klon Desktop Pro** | **Inkscape** | C++, GTK | Standar industri editor vektor berbasis SVG untuk ilustrasi & desain presisi. |
| **PWA Web Editor** | **Graphite** | Rust, WebAssembly | Editor 2D prosedural berbasis node yang berjalan di browser. |
| **Web SVG Editor** | **SVG-Edit & Method Draw** | JavaScript | Editor SVG cepat di browser tanpa instalasi. |
| **WebGPU Editor** | **Vectora** | JS, WGSL, WebGPU | Editor vektor masa depan berbasis WebGPU tanpa instalasi paket. |
| **UI/UX Platform** | **Penpot** | Clojure, TypeScript | Platform desain antarmuka open-source berbasis web dengan standar SVG. |
| **AI Foundation** | **SVGThinker-7B** | LLM | Model penalaran kecerdasan buatan khusus untuk menghasilkan kode SVG kompleks. |
| **AI Foundation** | **StarVector-8B** | AI Model | Mengubah gambar raster ke instruksi vektor SVG presisi tinggi. |
| **Render Engine** | **ThorVG** | C++, WebGPU | Mesin gambar vektor berkinerja tinggi untuk SVG dan animasi Lottie. |
| **Konverter Vektor** | **VTracer & Potrace** | Rust / C | Konversi raster ke vektor SVG full-color & centerline line art. |
    """)

# -----------------------------------------------------------------------------
# Footer
# -----------------------------------------------------------------------------
st.markdown("---")
st.markdown("""
<div style="text-align: center; color: #64748b; font-size: 0.85rem; padding: 12px 0;">
    <b>CorelDoc Studio Pro</b> &bull; Dibangun untuk Kompatibilitas Streamlit Cloud & GitHub Deployment &bull; 
    Repositori: <a href="https://github.com/indri007/corel" target="_blank" style="color: #00897B; text-decoration: none; font-weight: 600;">github.com/indri007/corel</a>
</div>
""", unsafe_allow_html=True)
